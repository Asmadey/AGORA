"""Извлечение текста из приложенного файла контекста аудитории.

Двоичные форматы разбираются в ВОРКЕРЕ намеренно. §6 CLAUDE.md запрещает
нативные npm-модули в apps/web, а §9 селит тяжёлые зависимости в образ
воркера — иначе появляется вторая среда, которая молча разойдётся с первой.
Здесь же разбор соседствует с остальным конвейером, поэтому отсутствие
библиотеки обнаруживается ДО оплаченной генерации, а не вместо неё.
"""
from __future__ import annotations

import re
from pathlib import Path

CONTEXT_LIMIT_CHARS = 4000

# Контекст попадает в системный промпт каждой персоны. На боевом прогоне 0051
# базовый промпт был около 83 000 символов (примерно 36 000 токенов); для
# русского текста 4000 символов — около 1500 токенов. При 27 персонах и
# перекрытии ×3 это около 121 000 токенов, то есть примерно 4% от 2,8 млн
# токенов прогона. Потолок ограничивает повторяемую цену, а не окно модели.


class ContextExtractionError(ValueError):
    """Файл нельзя превратить в непустой текст для portrait.distill."""


def _normalize(text: str) -> str:
    return re.sub(r"\n{3,}", "\n\n", text.replace("\r\n", "\n").replace("\r", "\n")).strip()


def _pdf_text(path: Path) -> str:
    try:
        from pdfminer.high_level import extract_text
    except ImportError as exc:  # pragma: no cover - worker image owns the dependency
        raise ContextExtractionError(
            "pdfminer.six не установлен в образе воркера"
        ) from exc

    try:
        return extract_text(str(path))
    except Exception as exc:  # noqa: BLE001 - expose a user-facing failed reason
        raise ContextExtractionError(f"PDF не удалось прочитать: {exc}") from exc


def _spreadsheet_text(path: Path) -> str:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - worker image owns the dependency
        raise ContextExtractionError("openpyxl не установлен в образе воркера") from exc

    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 - include the actual format failure
        raise ContextExtractionError(f"таблицу Excel не удалось прочитать: {exc}") from exc

    lines: list[str] = []
    try:
        for sheet in workbook.worksheets:
            for row in sheet.iter_rows(values_only=True):
                cells = [str(value).strip() for value in row if value is not None]
                if cells:
                    lines.append("\t".join(cells))
    finally:
        workbook.close()
    return "\n".join(lines)


def extract_context_text(path: str | Path) -> str:
    """Отдаёт не больше 4000 нормализованных символов либо падает громко.

    Потолок применяется ПОСЛЕ извлечения: размер PDF и таблицы ничего не
    говорит о количестве полезного текста — тяжёлый скан даёт ноль слов, а
    лёгкая таблица может дать десять тысяч.

    Скан без текстового слоя — это НЕУДАВШИЙСЯ контекст, а не удавшийся
    пустой. Разница принципиальная: пустая строка уехала бы в промпт как
    «контекста нет», и человек до конца прогона считал бы, что его файл учтён.
    """
    source = Path(path)
    suffix = source.suffix.lower()
    if suffix == ".pdf":
        raw = _pdf_text(source)
    elif suffix == ".xls":
        # openpyxl читает только OOXML. Старый BIFF-контейнер ему не по зубам,
        # и без этой ветки он падал бы невнятным InvalidFileException уже
        # внутри оплаченного прогона. Веб отклоняет `.xls` раньше, при выборе
        # файла (lib/context-file.ts), но путь в воркер существует и помимо
        # браузера — поэтому причина названа и здесь.
        raise ContextExtractionError(
            "формат .xls не читается: пересохраните таблицу как .xlsx"
        )
    elif suffix == ".xlsx":
        raw = _spreadsheet_text(source)
    else:
        raise ContextExtractionError(
            f"формат {suffix or 'без расширения'} не поддерживается для контекста"
        )

    text = _normalize(raw)
    if not text:
        raise ContextExtractionError(
            "из файла не извлечено ни одного слова: PDF может быть сканом без текстового слоя"
        )
    return text[:CONTEXT_LIMIT_CHARS]
