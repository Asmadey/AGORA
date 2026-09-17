"""
Выгрузка ответов в Excel — в форме полевого файла заказчика.

─── Почему в воркере, а не в вебе ───────────────────────────────────────────
§6 запрещает нативные npm-модули в `apps/web`, и это ровно тот случай: любая
сборка xlsx на стороне Node — либо нативный пакет, либо тяжёлый чистый JS,
который придётся тащить в клиентский бандл. В воркере уже есть `openpyxl`, а
готовый файл уезжает существующей подписанной ссылкой.

─── Откуда взята форма ──────────────────────────────────────────────────────
Из файла заказчика `Тюремный_дневник_Заполненные_респондентами_анкеты.xlsx`
(28 респондентов, 183 колонки), разобранного 17.09.2026:

    строка 1  блок, объединён по своим колонкам («Оценки 1-й серии»)
    строка 2  формулировка вопроса, объединена по колонкам вопроса
    строка 3+ респондент

    A        B              C     D
    Населенный пункт | ID | Пол | Возраст | …вопросы…

    E … I    «Какие ценности являются наиболее важными?»
             Права и свободы | Ответственность за себя | … | (пусто)

Ключевое, что легко сделать иначе и молча испортить: множественный выбор
занимает столько соседних колонок, сколько разрешено ответов, и в ячейке лежит
ТЕКСТ варианта, а не 0/1. Флаговая матрица читается инструментами заказчика как
другой файл — сравнить волны колонка в колонку станет нельзя.

─── Строка — персона ────────────────────────────────────────────────────────
Решение владельца 17.09.2026. При перекрытии больше единицы числовые поля
усредняются по повторам, категориальные берутся по моде. Расхождение с отчётом
при этом возможно: агрегат считает каждый повтор отдельным наблюдением. Поэтому
лист «О прогоне» называет и число персон, и число ответов — молча разойтись они
не должны.
"""

from __future__ import annotations

import statistics
from collections import Counter
from typing import Any

from ..survey import parse_field_answer, question_options, question_rows, survey_questions

#: Параметры аудитории, стоящие слева до вопросов. Пары «подпись → поле DNA».
#:
#: Состав и порядок повторяют полевой файл заказчика; `geo` и `children`
#: добавлены потому, что в его других выгрузках они есть, а у нас они всегда
#: заполнены.
AUDIENCE_COLUMNS: tuple[tuple[str, str], ...] = (
    ("Населенный пункт", "city"),
    ("ID респондента", "id"),
    ("Пол", "gender"),
    ("Возраст", "age"),
    ("Тип населённого пункта", "geo"),
    ("Наличие детей", "children"),
)

HEADER_ROWS = 2


def _demographics(persona: dict[str, Any]) -> dict[str, Any]:
    return (persona.get("dna") or {}).get("demographics") or {}


def _audience_value(persona: dict[str, Any], key: str) -> Any:
    if key == "id":
        return persona.get("id")
    return _demographics(persona).get(key)


def _answers_of(answer: dict[str, Any]) -> dict[str, Any]:
    body = answer.get("answer") if isinstance(answer.get("answer"), dict) else answer
    raw = body.get("survey_answers")
    if isinstance(raw, dict):
        return dict(raw)
    out: dict[str, Any] = {}
    for pair in raw or []:
        if isinstance(pair, dict):
            key = str(pair.get("question") or "").strip()
            if key:
                out[key] = pair.get("answer")
    return out


def _collapse(values: list[Any]) -> Any:
    """
    Повторы одной персоны в одно значение.

    Числа усредняются, остальное берётся по моде. Первое встретившееся при
    ничьей — не «случайный выбор», а воспроизводимый: порядок повторов
    детерминирован.
    """
    present = [v for v in values if v is not None and v != ""]
    if not present:
        return None
    if all(isinstance(v, (int, float)) and not isinstance(v, bool) for v in present):
        mean = statistics.mean(present)
        return int(round(mean)) if all(isinstance(v, int) for v in present) else round(mean, 2)
    return Counter(map(str, present)).most_common(1)[0][0]


def block_labels(survey: Any) -> dict[str, str]:
    """
    Подписи блоков. Пусто — в шапку пойдут идентификаторы, и это будет видно.

    Заказчик читает верхнюю строку глазами: там стоит «Оценки сериала в целом»,
    а не `b1`. Идентификатор в шапке — не косметический промах: две выгрузки
    рядом перестают читаться, а ради этого форма и повторяется.
    """
    if isinstance(survey, dict):
        blocks = survey.get("blocks")
        if isinstance(blocks, list):
            return {
                str(b.get("id")): str(b.get("label") or b.get("id"))
                for b in blocks if isinstance(b, dict)
            }
    return {}


def _columns(
    questions: list[dict[str, Any]], labels: dict[str, str] | None = None
) -> list[dict[str, Any]]:
    """
    Развёртка анкеты в колонки.

    Одна колонка на поле, кроме мультивыбора: там их столько, сколько
    разрешено ответов. У матрицы колонка на строку, и подписью идёт текст
    строки — именно он стоит в файле заказчика, а не формулировка вопроса.
    """
    names = labels or {}
    out: list[dict[str, Any]] = []
    for q in questions:
        qtype = str(q.get("type") or "open")
        block_id = str(q.get("block") or "")
        block = names.get(block_id, block_id)
        if qtype == "matrix_single":
            for row in question_rows(q):
                out.append({
                    "block": block,
                    "label": str(row.get("label") or ""),
                    "field": str(row.get("id")),
                    "question": q,
                    "slot": 0,
                })
            continue
        slots = int(q.get("maxChoices") or 1) if qtype == "multi_choice" else 1
        for slot in range(slots):
            out.append({
                "block": block,
                "label": question_label_of(q),
                "field": str(q.get("id")),
                "question": q,
                "slot": slot,
            })
    return out


def question_label_of(question: dict[str, Any]) -> str:
    return str(question.get("label") or question.get("text") or "").strip()


def _cell(question: dict[str, Any], raw: Any, slot: int) -> Any:
    """
    Значение одной ячейки.

    Пропуск остаётся ПУСТЫМ, а не нулём: ноль — законный балл на шкале 0–10, и
    записать им «не ответил» значило бы утянуть среднее вниз на величину,
    которой никто не называл.
    """
    parsed = parse_field_answer(question, raw)
    qtype = str(question.get("type") or "open")

    if qtype == "scale":
        return parsed.value
    if qtype == "open":
        return parsed.text or None

    labels = {str(o.get("id")): str(o.get("label")) for o in question_options(question)}
    picked = [labels.get(oid, oid) for oid in parsed.option_ids]
    return picked[slot] if slot < len(picked) else None


def build_workbook(
    questions: list[dict[str, Any]] | dict[str, Any],
    answers: list[dict[str, Any]],
    personas: list[dict[str, Any]],
    *,
    meta: dict[str, Any] | None = None,
    blocks: dict[str, str] | None = None,
) -> Any:
    """
    Книга из двух листов: «Ответы» и «О прогоне».

    `blocks` — подписи блоков для верхней строки шапки. Если не переданы, они
    берутся из самого документа анкеты, когда он пришёл целиком; иначе в шапке
    останутся идентификаторы, и это будет видно с первого взгляда.
    """
    from openpyxl import Workbook
    from openpyxl.utils import get_column_letter

    qs = survey_questions(questions)
    columns = _columns(qs, blocks or block_labels(questions))

    by_persona: dict[str, list[dict[str, Any]]] = {}
    for a in answers:
        by_persona.setdefault(str(a.get("persona_id")), []).append(a)

    wb = Workbook()
    ws = wb.active
    ws.title = "Ответы"

    offset = len(AUDIENCE_COLUMNS)
    for i, (label, _) in enumerate(AUDIENCE_COLUMNS, start=1):
        ws.cell(row=2, column=i, value=label)

    for j, column in enumerate(columns, start=offset + 1):
        ws.cell(row=1, column=j, value=column["block"])
        ws.cell(row=2, column=j, value=column["label"])

    # Объединения по блоку и по вопросу — верхняя строка файла заказчика
    # объединена по всему блоку, вторая по колонкам одного вопроса.
    for key, row_no in (("block", 1), ("label", 2)):
        start = 0
        while start < len(columns):
            end = start
            while (
                end + 1 < len(columns)
                and columns[end + 1][key] == columns[start][key]
                and (key == "block" or columns[end + 1]["field"] == columns[start]["field"])
            ):
                end += 1
            if end > start and columns[start][key]:
                ws.merge_cells(
                    start_row=row_no, start_column=offset + 1 + start,
                    end_row=row_no, end_column=offset + 1 + end,
                )
            start = end + 1

    for r, persona in enumerate(personas, start=HEADER_ROWS + 1):
        for i, (_, key) in enumerate(AUDIENCE_COLUMNS, start=1):
            ws.cell(row=r, column=i, value=_audience_value(persona, key))

        replications = by_persona.get(str(persona.get("id")), [])
        raw_by_field: dict[str, list[Any]] = {}
        for a in replications:
            for field, value in _answers_of(a).items():
                raw_by_field.setdefault(field, []).append(value)

        for j, column in enumerate(columns, start=offset + 1):
            raw = _collapse(raw_by_field.get(column["field"], []))
            ws.cell(row=r, column=j, value=_cell(column["question"], raw, column["slot"]))

    ws.freeze_panes = ws.cell(row=HEADER_ROWS + 1, column=offset + 1)
    for i in range(1, offset + 1):
        ws.column_dimensions[get_column_letter(i)].width = 18

    about = wb.create_sheet("О прогоне")
    info = meta or {}
    rows = [
        ("Исследование", info.get("number") or info.get("run_id") or "—"),
        ("Дата выгрузки", info.get("exported_at") or "—"),
        ("Персон", len(personas)),
        ("Ответов", len(answers)),
        ("Вопросов в анкете", len(qs)),
        ("Полей к ответу", len(columns)),
        ("Версия анкеты", info.get("survey_version") or "—"),
        ("Модели", info.get("models") or "—"),
        ("Исключено правилами QA", info.get("excluded_by_qa", "—")),
        (
            "Как читать",
            "Строка — персона. При перекрытии больше единицы числа усреднены по "
            "повторам, категории взяты по моде; поэтому число ответов больше "
            "числа строк, а агрегат отчёта считает каждый повтор отдельно.",
        ),
    ]
    for r, (key, value) in enumerate(rows, start=1):
        about.cell(row=r, column=1, value=key)
        about.cell(row=r, column=2, value=value)
    about.column_dimensions["A"].width = 28
    about.column_dimensions["B"].width = 80
    return wb
